"""Unified, tolerant file I/O shared by every tool.

The point of this module is the thing the user asked for: one way to *read*
tabular data regardless of where it came from, and one way to *write* each
export format so the buttons behave identically across tools.

``read_table`` accepts .csv / .txt / .xlsx / .xls, auto-detects the delimiter
(tab, comma, semicolon or whitespace), recovers from encoding problems, and
finds the header row even when metadata precedes it. Tools that need
domain-specific extraction (unit scaling, cycle detection) build on top of
this or on the specialised readers in ``core.data_io``.
"""
from __future__ import annotations

import csv
import io as _io
import importlib.util
import zipfile
from dataclasses import dataclass
from pathlib import Path

import pandas as pd

_ENCODINGS = ("utf-8-sig", "utf-8", "cp1252", "latin-1")
TABLE_EXTS = (".csv", ".txt", ".tsv", ".xlsx", ".xls", ".dat")

# The readers intentionally have finite resource budgets.  These limits are
# well above normal electrochemistry/FTIR files while preventing an accidental
# multi-gigabyte export or a compressed XLSX bomb from exhausting the process.
MAX_TEXT_BYTES = 64 * 1024 * 1024
MAX_SPREADSHEET_BYTES = 100 * 1024 * 1024
MAX_XLSX_ENTRIES = 10_000
MAX_XLSX_UNCOMPRESSED_BYTES = 512 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 500
XLSX_RATIO_CHECK_BYTES = 32 * 1024 * 1024
MAX_TABLE_ROWS = 500_000
MAX_TABLE_CELLS = 5_000_000

_INVALID_SHEET_CHARS = frozenset("[]:*?/\\")
_FORMULA_PREFIXES = ("=", "+", "-", "@")


@dataclass(frozen=True)
class XlsxCellValue:
    """A workbook cell plus the type that openpyxl assigned to it.

    Keeping ``data_type`` is important: the string ``"=A1+1"`` can be either
    an intentional formula or literal text.  A bare Python string cannot
    represent that distinction when a source workbook is copied.
    """

    value: object
    data_type: str
    number_format: str = "General"


def _file_size(path: Path, limit: int, kind: str) -> int:
    try:
        size = path.stat().st_size
    except OSError as exc:
        raise ValueError(f"Could not access {kind} file '{path.name}': {exc}") from exc
    if size > limit:
        raise ValueError(
            f"{kind.capitalize()} file '{path.name}' is too large "
            f"({size / (1024 * 1024):.1f} MiB; limit {limit / (1024 * 1024):.0f} MiB)."
        )
    return size


def _guard_xlsx_archive(path: Path) -> None:
    """Reject oversized or suspicious XLSX zip containers before XML parsing."""
    _file_size(path, MAX_SPREADSHEET_BYTES, "spreadsheet")
    try:
        with zipfile.ZipFile(path) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_ENTRIES:
                raise ValueError(
                    f"Spreadsheet '{path.name}' contains too many archive entries "
                    f"({len(members):,}; limit {MAX_XLSX_ENTRIES:,})."
                )
            total_raw = sum(item.file_size for item in members)
            total_packed = sum(item.compress_size for item in members)
            if total_raw > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValueError(
                    f"Spreadsheet '{path.name}' expands to {total_raw / (1024 * 1024):.1f} MiB "
                    f"(limit {MAX_XLSX_UNCOMPRESSED_BYTES / (1024 * 1024):.0f} MiB)."
                )
            ratio = total_raw / max(total_packed, 1)
            if total_raw >= XLSX_RATIO_CHECK_BYTES and ratio > MAX_XLSX_COMPRESSION_RATIO:
                raise ValueError(
                    f"Spreadsheet '{path.name}' has a suspicious decompression ratio "
                    f"({ratio:.0f}:1; limit {MAX_XLSX_COMPRESSION_RATIO}:1)."
                )
    except zipfile.BadZipFile as exc:
        raise ValueError(f"Spreadsheet '{path.name}' is not a valid .xlsx file.") from exc


def _guard_shape(rows: int, columns: int, label: str) -> None:
    cells = rows * columns
    if rows > MAX_TABLE_ROWS:
        raise ValueError(
            f"{label} has too many rows ({rows:,}; limit {MAX_TABLE_ROWS:,})."
        )
    if cells > MAX_TABLE_CELLS:
        raise ValueError(
            f"{label} has too many cells ({cells:,}; limit {MAX_TABLE_CELLS:,})."
        )


def _spreadsheet_engine(path: Path) -> str:
    if path.suffix.lower() == ".xls":
        if importlib.util.find_spec("xlrd") is None:
            raise ValueError(
                "Legacy .xls support requires xlrd>=2.0.1. "
                "Install the application's requirements or convert the file to .xlsx."
            )
        return "xlrd"
    return "openpyxl"


def read_excel(path, **kwargs) -> pd.DataFrame:
    """Guarded pandas Excel reader with explicit engines and clear .xls errors."""
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        _guard_xlsx_archive(path)
    else:
        _file_size(path, MAX_SPREADSHEET_BYTES, "spreadsheet")
    frame = pd.read_excel(path, engine=_spreadsheet_engine(path), **kwargs)
    _guard_shape(len(frame), len(frame.columns), f"Spreadsheet '{path.name}'")
    return frame


def _read_text(path: Path) -> str:
    """Read a bounded text file, trying the supported encodings in order."""
    _file_size(path, MAX_TEXT_BYTES, "text")
    for enc in _ENCODINGS:
        try:
            with open(path, "r", encoding=enc) as fh:
                return fh.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    with open(path, "r", encoding="utf-8", errors="replace") as fh:
        return fh.read()


def read_text(path) -> str:
    """Read a text file with shared encoding fallbacks and the size budget."""
    return _read_text(Path(path))


def _sniff_delimiter(sample: str) -> str:
    """Best-effort delimiter detection for a text sample."""
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters="\t,; ")
        return dialect.delimiter
    except csv.Error:
        pass
    # Manual fallback: pick the candidate that appears most on the first
    # data-looking line.
    for line in sample.splitlines():
        if line.strip():
            counts = {d: line.count(d) for d in ("\t", ",", ";")}
            best = max(counts, key=counts.get)
            if counts[best] > 0:
                return best
            break
    return r"\s+"


def _looks_numeric(token: str) -> bool:
    try:
        float(token)
        return True
    except ValueError:
        return False


def _find_header_row(lines: list[str], delim: str) -> int:
    """Return the index of the first row that looks like a header (mostly
    non-numeric tokens, at least two columns), else 0."""
    for idx, line in enumerate(lines[:60]):
        if delim in {" ", r"\s+"}:
            row = line.split()
        else:
            try:
                row = next(csv.reader([line], delimiter=delim))
            except csv.Error:
                row = line.split(delim)
        toks = [t.strip() for t in row if t.strip() != ""]
        if len(toks) < 2:
            continue
        non_numeric = sum(0 if _looks_numeric(t) else 1 for t in toks)
        if non_numeric >= max(1, len(toks) // 2):
            return idx
    return 0


def read_table(path, header: str = "auto") -> pd.DataFrame:
    """Read one tabular file into a DataFrame as tolerantly as possible.

    header: "auto" (detect a header row), "yes" (row 0 is header) or
    "no" (no header — columns become 0,1,2,...).
    """
    path = Path(path)
    ext = path.suffix.lower()

    if ext in (".xlsx", ".xls"):
        hdr = 0 if header != "no" else None
        return read_excel(path, header=hdr)

    text = _read_text(path)
    lines = [ln for ln in text.splitlines() if ln.strip() != ""]
    if not lines:
        return pd.DataFrame()

    delim = "," if ext == ".csv" else _sniff_delimiter("\n".join(lines[:40]))

    if header == "no":
        skip, hdr = 0, None
    elif header == "yes":
        skip, hdr = 0, 0
    else:
        skip, hdr = _find_header_row(lines, delim), 0

    # ``csv.Sniffer`` can return a literal space.  Treat it as generic
    # whitespace so repeated spaces do not create empty phantom columns.
    sep = r"\s+" if delim in {" ", r"\s+"} else delim
    try:
        df = pd.read_csv(
            _io.StringIO(text), sep=sep, skiprows=skip, header=hdr,
            engine="python", on_bad_lines="skip",
        )
    except Exception:
        df = pd.read_csv(
            _io.StringIO(text), sep=r"\s+", header=None,
            engine="python", on_bad_lines="skip",
        )
    _guard_shape(len(df), len(df.columns), f"Table '{path.name}'")
    return df


def read_text_grid(path) -> list[list[str]]:
    """Read a tab/whitespace-delimited text file into a raw grid of string
    cells, preserving every row (including headers and non-numeric text).

    Used by the TXT->XLSX tool, which must keep the file's literal layout
    rather than coercing it into a typed DataFrame.
    """
    path = Path(path)
    text = _read_text(path)
    grid: list[list[str]] = []
    ext = path.suffix.lower()
    sample = "\n".join(text.splitlines()[:40])
    delim = "," if ext == ".csv" else _sniff_delimiter(sample)

    if delim in {" ", r"\s+"}:
        records = (line.split() for line in text.splitlines())
    else:
        # csv.reader handles escaped quotes, delimiters inside quotes, and
        # quoted fields spanning physical lines.  It is also correct for TSV
        # and semicolon-separated files when given the detected delimiter.
        records = csv.reader(_io.StringIO(text), delimiter=delim)

    cell_count = 0
    try:
        for record in records:
            cells = list(record)
            if not cells or not any(str(cell).strip() for cell in cells):
                continue
            grid.append(cells)
            cell_count += len(cells)
            if len(grid) > MAX_TABLE_ROWS:
                raise ValueError(
                    f"Table '{path.name}' has too many rows "
                    f"(limit {MAX_TABLE_ROWS:,})."
                )
            if cell_count > MAX_TABLE_CELLS:
                raise ValueError(
                    f"Table '{path.name}' has too many cells "
                    f"(limit {MAX_TABLE_CELLS:,})."
                )
    except csv.Error as exc:
        raise ValueError(f"Could not parse delimited file '{path.name}': {exc}") from exc
    return grid


def read_xlsx_grids(path) -> list[tuple[str, list[list[XlsxCellValue | str]]]]:
    """Read an .xlsx file into one typed grid per worksheet.

    Returns ``[(sheet_title, grid), ...]`` in workbook order, where each grid is
    a list of rows of :class:`XlsxCellValue` objects.  Formula cells retain
    their formulas (``data_only=False``), and their original data type remains
    distinguishable from literal text beginning with ``=``. ``None`` cells
    become ``""``.

    Used by the TXT->XLSX tool to expand each input workbook's sheets. Invalid
    or unsafe workbooks raise a user-facing ``ValueError`` rather than silently
    disappearing from the output.
    """
    from openpyxl import load_workbook

    path = Path(path)
    _guard_xlsx_archive(path)
    try:
        wb = load_workbook(
            path, read_only=True, data_only=False, keep_links=False,
        )
    except Exception as exc:
        raise ValueError(f"Could not read spreadsheet '{path.name}': {exc}") from exc

    grids: list[tuple[str, list[list[XlsxCellValue | str]]]] = []
    total_cells = 0
    try:
        for ws in wb.worksheets:
            _guard_shape(
                ws.max_row or 0,
                ws.max_column or 0,
                f"Worksheet '{ws.title}' in '{path.name}'",
            )
            rows: list[list[XlsxCellValue | str]] = []
            for row in ws.iter_rows(values_only=False):
                converted: list[XlsxCellValue | str] = []
                for cell in row:
                    if cell.value is None:
                        converted.append("")
                    else:
                        converted.append(XlsxCellValue(
                            cell.value, cell.data_type, cell.number_format,
                        ))
                rows.append(converted)
                total_cells += len(converted)
                if len(rows) > MAX_TABLE_ROWS or total_cells > MAX_TABLE_CELLS:
                    raise ValueError(
                        f"Spreadsheet '{path.name}' exceeds the table limit "
                        f"({MAX_TABLE_ROWS:,} rows or {MAX_TABLE_CELLS:,} cells)."
                    )
            # Drop fully-empty trailing rows.
            while rows and all(c == "" for c in rows[-1]):
                rows.pop()
            grids.append((ws.title, rows))
    finally:
        wb.close()
    return grids


# ---------------------------------------------------------------------------
# Writers — one per export kind, used by widgets.export_button
# ---------------------------------------------------------------------------
def neutralize_spreadsheet_formula(value):
    """Return formula-looking user text as inert text for CSV/XLSX exports.

    Spreadsheet applications recognize ``=``, ``+``, ``-`` and ``@`` as
    formula prefixes.  A leading apostrophe is the portable way to force a
    value to remain text.  Leading whitespace is ignored only for detection so
    an attacker cannot evade the guard with spaces or control characters.
    """
    if not isinstance(value, str) or value.startswith("'"):
        return value
    candidate = value.lstrip(" \t\r\n\ufeff")
    if candidate.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _sanitized_dataframe(df: pd.DataFrame) -> pd.DataFrame:
    if not isinstance(df, pd.DataFrame):
        raise TypeError("Export payload must be a pandas DataFrame.")
    clean = df.apply(lambda column: column.map(neutralize_spreadsheet_formula))
    clean.columns = [neutralize_spreadsheet_formula(str(c)) for c in df.columns]
    return clean


def safe_sheet_title(value, used_names: set[str] | None = None) -> str:
    """Return a valid, <=31-character, case-insensitively unique Excel title."""
    base = "".join(
        "_" if char in _INVALID_SHEET_CHARS or ord(char) < 32 else char
        for char in str(value)
    ).strip()
    base = base.strip("'")[:31] or "Sheet"
    used_folded = {name.casefold() for name in (used_names or set())}
    candidate = base
    counter = 1
    while candidate.casefold() in used_folded:
        suffix = f"_{counter}"
        candidate = base[:31 - len(suffix)] + suffix
        counter += 1
    if used_names is not None:
        used_names.add(candidate)
    return candidate


def write_csv(df: pd.DataFrame, path) -> None:
    # A UTF-8 BOM makes Excel and other Windows spreadsheet programs detect
    # the encoding instead of treating multi-byte unit symbols (Ω, ·, ², µ,
    # etc.) as characters from the active legacy code page.
    _sanitized_dataframe(df).to_csv(path, index=False, encoding="utf-8-sig")


def write_xlsx(sheets: dict, path, number_format: str | None = None) -> None:
    """Write formatted result tables with frozen headers to one .xlsx."""
    from workbench.core.excel import format_sheet

    used_names: set[str] = set()
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, df in sheets.items():
            safe = safe_sheet_title(name, used_names)
            clean = _sanitized_dataframe(df)
            _guard_shape(len(clean) + 1, len(clean.columns), f"Export sheet '{safe}'")
            clean.to_excel(writer, sheet_name=safe, index=False)
            format_sheet(writer.sheets[safe], number_format=number_format)


def write_txt(text: str, path) -> None:
    # Keep standalone text exports consistent with CSV exports so Windows
    # applications preserve scientific symbols when opening the file directly.
    with open(path, "w", encoding="utf-8-sig") as fh:
        fh.write(text)


def save_figure(fig, path) -> None:
    fig.savefig(path, dpi=200, bbox_inches="tight")


def expand_inputs(paths, accepted_exts) -> list[Path]:
    """Expand a mixed list of files and directories into a flat list of files
    whose extension is in ``accepted_exts``. Directories are searched one level
    deep (then recursively if nothing matched at the top)."""
    accepted = tuple(e.lower() for e in accepted_exts)
    out: list[Path] = []
    for p in paths:
        p = Path(p)
        if p.is_dir():
            matched = [f for f in sorted(p.iterdir())
                       if f.is_file() and f.suffix.lower() in accepted]
            if not matched:
                matched = [f for f in sorted(p.rglob("*"))
                           if f.is_file() and f.suffix.lower() in accepted]
            out.extend(matched)
        elif p.is_file() and (not accepted or p.suffix.lower() in accepted):
            out.append(p)
    return out
