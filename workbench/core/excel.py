"""Readable worksheet formatting for analysis and conversion exports."""
from __future__ import annotations

import math
from numbers import Real
import textwrap
import unicodedata

from openpyxl.styles import Alignment, Border, Font, PatternFill


_BODY = Font(name="Calibri", size=11, color="000000")
_HEADER = Font(name="Calibri", size=11, bold=True, color="000000")
_NO_FILL = PatternFill()
_NO_BORDER = Border()
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_NUMBER_ALIGNMENT = Alignment(horizontal="right", vertical="center")
_TEXT_ALIGNMENT = Alignment(horizontal="left", vertical="center")
_WRAPPED_ALIGNMENT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def _width(text: str) -> int:
    return max((sum(2 if unicodedata.east_asian_width(char) in "WF" else 1
                    for char in line) for line in text.splitlines()), default=0)


def _number_format(value: Real) -> str:
    if value != 0 and (abs(value) < 1e-4 or abs(value) >= 1e7):
        return "0.000E+00"
    return "#,##0" if isinstance(value, int) else "0.0#####"


def _display_text(cell) -> str:
    value = cell.value
    if isinstance(value, Real) and not isinstance(value, bool):
        if "E+00" in cell.number_format:
            return f"{value:.3E}"
        if cell.number_format == "#,##0":
            return f"{value:,.0f}"
        if cell.number_format in {"0.0#####", "0.######"}:
            text = f"{value:.6f}".rstrip("0").rstrip(".")
            return text + ".0" if cell.number_format == "0.0#####" and "." not in text else text
    return str(value)


def _line_count(text: str, width: float) -> int:
    return sum(max(1, len(textwrap.wrap(line, width=int(width))),
                   math.ceil(_width(line) / width))
               for line in text.splitlines() or [""])


def format_sheet(sheet, *, header_row: int | None = 1,
                 header_rows: tuple[int, ...] | None = None,
                 number_format: str | None = None,
                 preserve_number_formats: bool = False) -> None:
    """Use a plain Excel grid and fit content without changing cell values."""
    sheet.sheet_view.showGridLines = True
    headers = set(header_rows if header_rows is not None else
                  (() if header_row is None else (header_row,)))
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            cell.font = _HEADER if cell.row in headers else _BODY
            cell.fill = _NO_FILL
            cell.border = _NO_BORDER
            numeric = isinstance(cell.value, Real) and not isinstance(cell.value, bool)
            if (numeric and cell.row not in headers and not preserve_number_formats
                    and math.isfinite(cell.value)):
                cell.number_format = number_format or _number_format(cell.value)
    widths = {}
    for column in range(1, sheet.max_column + 1):
        samples = [sheet.cell(row, column) for row in range(1, min(sheet.max_row, 300) + 1)]
        text_width = 0
        for cell in samples:
            if cell.value is None or cell.data_type == "f":
                continue
            content_width = _width(_display_text(cell))
            if cell.row in headers:
                content_width = min(32, content_width)
            elif isinstance(cell.value, str):
                content_width = math.ceil(content_width * 1.12)
            text_width = max(text_width, content_width)
        width = min(52, max(10, text_width + 2))
        widths[column] = width
        sheet.column_dimensions[sheet.cell(1, column).column_letter].width = width
    if header_row is not None:
        sheet.freeze_panes = f"A{header_row + 1}"
    else:
        sheet.freeze_panes = None
    for row in sheet.iter_rows():
        row_height = 15
        for cell in row:
            value = cell.value
            if value is None:
                continue
            numeric = isinstance(value, Real) and not isinstance(value, bool)
            if cell.row in headers:
                cell.alignment = _HEADER_ALIGNMENT
                row_height = max(row_height, 15 * _line_count(str(value), widths[cell.column] - 2))
            elif numeric:
                cell.alignment = _NUMBER_ALIGNMENT
            else:
                wrapped = cell.data_type != "f" and (
                    "\n" in str(value) or _width(str(value)) > widths[cell.column] - 2
                )
                cell.alignment = _WRAPPED_ALIGNMENT if wrapped else _TEXT_ALIGNMENT
                if wrapped:
                    row_height = max(row_height, 15 * _line_count(str(value), widths[cell.column] - 2))
        if any(cell.value is not None for cell in row):
            sheet.row_dimensions[row[0].row].height = row_height
