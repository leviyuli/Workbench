"""Restored project plots and exports retain their original presentation/data."""

from io import BytesIO

import pandas as pd
import pytest
from matplotlib.figure import Figure
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from PySide6.QtCore import QCoreApplication, QEvent
from PySide6.QtWidgets import QApplication
from shiboken6 import isValid

from workbench.core.tool_base import Artifact, Tool, ToolResult
from workbench.widgets.export_button import ExportButton
from workbench.widgets.figure_view import SingleFigure, StaticFigure
from workbench.widgets.results_panel import ResultsPanel


@pytest.fixture(scope="module")
def qt_app():
    app = QApplication.instance()
    if app is None:
        with pytest.MonkeyPatch.context() as monkeypatch:
            monkeypatch.setenv("QT_QPA_PLATFORM", "offscreen")
            app = QApplication([])
    yield app


@pytest.fixture
def png():
    figure = Figure(figsize=(4, 2))
    figure.subplots().plot([0, 1], [0, 2])
    buffer = BytesIO()
    figure.savefig(buffer, format="png")
    return buffer.getvalue()


@pytest.fixture
def panel(qt_app):
    widget = ResultsPanel()
    yield widget
    widget.deleteLater()
    QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)


@pytest.mark.parametrize("display_tables", [False, True])
@pytest.mark.parametrize("display_summary", [False, True])
def test_restored_plot_order_and_display_flags(panel, png, display_tables, display_summary):
    result = ToolResult(
        figure_images={"Electrodes": png, "Cell comparison": png},
        tables={"Source": pd.DataFrame({"Current": [0.0, -1.0]})},
        scalars={"Offset": -0.2}, text_blocks={"Alignment": "Cathode onset"},
        warnings=["Example warning"], display_tables=display_tables,
        display_summary=display_summary,
    )
    panel.set_result(result)
    names = [panel.tabs.tabText(index) for index in range(panel.tabs.count())]
    expected = ["Electrodes", "Cell comparison"]
    if display_summary:
        expected.append("Summary")
    if display_tables:
        expected.append("Source")
    assert names == expected
    assert isinstance(panel.tabs.widget(0), StaticFigure)
    assert not panel.findChildren(SingleFigure)
    assert "interactive" in panel.tabs.widget(0).hint.text()
    assert "Example warning" in panel.summary.toPlainText()
    assert result.tables["Source"].Current.tolist() == [0.0, -1.0]


def test_saved_image_scales_with_viewport_without_distortion(qt_app, png):
    view = StaticFigure(png)
    try:
        view.resize(700, 500)
        view.show()
        qt_app.processEvents()
        first = view.canvas.image_rect()
        view.resize(1000, 700)
        qt_app.processEvents()
        second = view.canvas.image_rect()
        assert second.width() > first.width()
        for rect in (first, second):
            assert abs(rect.width() / rect.height() - 2.0) < 0.02
        assert view.canvas.rect().contains(second)
        assert not hasattr(view.canvas, "figure")
    finally:
        view.deleteLater()


def test_rerunning_replaces_snapshot_with_live_figure(panel, png):
    panel.set_result(ToolResult(figure_images={"Plot": png}))
    old_snapshot = panel.tabs.widget(0)
    figure = Figure()
    figure.subplots().plot([0, 1], [0, 1])
    panel.set_result(ToolResult(figures={"Plot": figure}))
    QCoreApplication.sendPostedEvents(None, QEvent.DeferredDelete)
    assert not isValid(old_snapshot)
    assert isinstance(panel.tabs.widget(0), SingleFigure)
    assert panel.tabs.widget(0).canvas.figure is figure
    assert not panel.findChildren(StaticFigure)


def test_live_figure_takes_precedence_over_same_name_snapshot(panel, png):
    panel.set_result(ToolResult(figures={"Plot": Figure()}, figure_images={"Plot": png}))
    assert panel.tabs.count() == 2
    assert isinstance(panel.tabs.widget(0), SingleFigure)


def test_invalid_saved_image_has_clear_error(qt_app):
    with pytest.raises(ValueError, match="valid PNG"):
        StaticFigure(b"not an image")


@pytest.mark.parametrize("kind", ["csv", "xlsx", "workbook", "png", "txt"])
def test_serialized_exports_are_written_verbatim(kind, tmp_path):
    data = b"\xef\xbb\xbfOriginal\x00data\r\n"
    artifact = Artifact("Original export", kind, f"output.{kind}", object(), data)
    target = tmp_path / artifact.default_filename
    ExportButton._write(artifact, str(target))
    assert target.read_bytes() == data


def test_serialized_workbook_preserves_formulas_types_and_formatting(tmp_path):
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "=1+2"
    sheet["B1"] = "=literal text"
    sheet["B1"].data_type = "s"
    sheet["C1"] = 1.23456
    sheet["C1"].number_format = "0.0000"
    sheet["C1"].font = Font(bold=True)
    buffer = BytesIO()
    book.save(buffer)
    target = tmp_path / "restored.xlsx"
    ExportButton._write(Artifact("Workbook", "workbook", target.name, None,
                                 buffer.getvalue()), str(target))
    restored = load_workbook(target)
    try:
        output = restored.active
        assert output["A1"].data_type == "f"
        assert output["B1"].data_type == "s"
        assert output["B1"].value == "=literal text"
        assert output["C1"].value == 1.23456
        assert output["C1"].number_format == "0.0000"
        assert output["C1"].font.bold
    finally:
        restored.close()


def test_live_artifacts_still_use_normal_writers(tmp_path, png):
    table_path = tmp_path / "values.csv"
    ExportButton._write(Artifact("Data", "csv", table_path.name,
                                 pd.DataFrame({"Value": [1.23456789]})),
                        str(table_path))
    assert pd.read_csv(table_path).Value.iloc[0] == 1.23456789
    figure = Figure()
    figure.subplots().plot([0, 1], [0, 1])
    image_path = tmp_path / "figure.png"
    ExportButton._write(Artifact("Figure", "png", image_path.name, figure),
                        str(image_path))
    assert image_path.read_bytes().startswith(b"\x89PNG\r\n\x1a\n")


def test_default_artifacts_include_restored_png_without_duplicate(qt_app, png):
    class ExampleTool(Tool):
        def parameter_schema(self):
            return []

        def run(self, files, params, extra=None):
            return ToolResult()

    result = ToolResult(figures={"Live": Figure()},
                        figure_images={"Saved": png, "Live": png})
    artifacts = ExampleTool().exportable_artifacts(result)
    assert [artifact.name for artifact in artifacts] == ["Live (PNG)", "Saved (PNG)"]
    assert artifacts[0].serialized_payload is None
    assert artifacts[1].serialized_payload == png


def test_restored_export_directory_can_use_project_folder(qt_app, tmp_path):
    button = ExportButton()
    button.set_input_dir(str(tmp_path))
    assert button._start_dir() == str(tmp_path)
    button.deleteLater()
