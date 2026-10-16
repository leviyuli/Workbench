"""Integrated RE exports retain plotted series without analysis intermediates."""

from io import BytesIO

import numpy as np
from openpyxl import load_workbook
import pandas as pd
import pytest

from workbench.core.project import (
    artifact_bytes, load_project, restore_result, save_project, snapshot_result,
)
from workbench.tools.integrated_re.exports import plot_tables
from workbench.tools.integrated_re.tool import IntegratedReTool
from workbench.widgets.export_button import ExportButton


def _params(**changes):
    values = {spec.key: spec.default for spec in IntegratedReTool().parameter_schema()}
    values.update(calibration_reference="RHE", calibration_ph=0.0,
                  calibration_temperature=298.15)
    values.update(changes)
    return values


def _curve(path, potential, current):
    pd.DataFrame({"E(V)": potential, "j(A/cm2)": current}).to_csv(path, index=False)
    return str(path)


def _sources(folder, *, calibration=False, cell=True, gaps=False):
    folder.mkdir(parents=True, exist_ok=True)
    ja = np.array([-0.025, 0.0, 0.1, 0.3, 0.5])
    jc = np.array([0.02, 0.0, -0.15, -0.2, -0.35, -0.5])
    ea = 1.4123456789012 + 0.8 * ja
    if gaps:
        ea[2] = np.nan
    sources = {
        "anode": _curve(folder / "anode.csv", ea, ja),
        "cathode": _curve(folder / "cathode.csv", -0.1 + 0.4 * jc, jc),
        "eis": {},
    }
    for role, resistance in (("anode", 0.2), ("cathode", 0.3), ("cell", 0.7)):
        path = folder / f"{role}_eis.csv"
        path.write_text(
            "Frequency (Hz),Re(Z) (Ohm.cm2),-Im(Z) (Ohm.cm2)\n"
            f"10000,{resistance - 0.05},-0.1\n"
            f"1000,{resistance + 0.05},0.1\n"
            f"100,{resistance + 0.3},0.5\n",
            encoding="utf-8",
        )
        sources["eis"][role] = [str(path)]
    if cell:
        jcell = np.array([-0.04, 0.05, 0.2, 0.4, 0.6])
        sources["cell"] = _curve(folder / "cell.csv", 1.51 + 1.3 * jcell, jcell)
    if calibration:
        j3e = np.array([0.0, -0.1, -0.3, -0.5])
        sources["calibration"] = _curve(folder / "calibration.csv", -0.14 + 0.4 * j3e, j3e)
    return sources


def _line(result, figure, label, *, corrected=False):
    lines = [line for line in result.figures[figure].axes[0].lines
             if label.lower() in line.get_label().lower()
             and ("ir corrected" in line.get_label().lower()) == corrected]
    assert len(lines) == 1
    return lines[0]


def _matches_line(table, x_column, y_column, line):
    x, y = np.asarray(line.get_xdata()), np.asarray(line.get_ydata())
    np.testing.assert_allclose(table[x_column].iloc[:len(x)], x, equal_nan=True)
    np.testing.assert_allclose(table[y_column].iloc[:len(y)], y, equal_nan=True)
    assert table[x_column].iloc[len(x):].isna().all()
    assert table[y_column].iloc[len(y):].isna().all()


@pytest.mark.parametrize("axis", ["E vs RHE", "Overpotential"])
@pytest.mark.parametrize("include_ir", [False, True])
def test_plot_data_matches_electrode_and_cell_curves(tmp_path, axis, include_ir):
    tool = IntegratedReTool()
    result = tool.run([], _params(electrode_axis=axis, include_ir=include_ir,
                                 correction_mode="Manual", manual_shift=-0.1234567890123),
                      extra=_sources(tmp_path, gaps=True))
    tables = plot_tables(result)
    assert list(tables) == ["Cell comparison", "Electrode LSVs"]
    electrodes, cell = tables["Electrode LSVs"], tables["Cell comparison"]
    for role in ("Anode", "Cathode"):
        voltage = f"{role} E (V vs RHE)" if axis == "E vs RHE" else f"{role} η (V)"
        current = f"{role} j (A/cm²)"
        _matches_line(electrodes, voltage, current,
                      _line(result, "Electrode LSVs", role))
        corrected_voltage = (f"{role} E iR corrected (V vs RHE)" if axis == "E vs RHE"
                             else f"{role} η iR corrected (V)")
        if include_ir:
            _matches_line(electrodes, corrected_voltage, current,
                          _line(result, "Electrode LSVs", role, corrected=True))
        else:
            assert corrected_voltage not in electrodes
    for label, prefix in (("Reconstructed", "Reconstructed"), ("Measured cell", "Measured (2E)")):
        _matches_line(cell, f"{prefix} V (V)", f"{prefix} |j| (A/cm²)",
                      _line(result, "Cell comparison", label))
        corrected_voltage = f"{prefix} V iR corrected (V)"
        if include_ir:
            _matches_line(cell, corrected_voltage, f"{prefix} |j| (A/cm²)",
                          _line(result, "Cell comparison", label, corrected=True))
        else:
            assert corrected_voltage not in cell
    assert np.isnan(electrodes["Anode j (A/cm²)"].iloc[2])
    assert (electrodes["Cathode j (A/cm²)"].dropna() < 0).any()
    assert electrodes["Anode j (A/cm²)"].iloc[-1:].isna().all()
    assert len(electrodes) == 6


def test_measured_source_is_exported_when_reconstruction_has_no_overlap(tmp_path):
    sources = _sources(tmp_path)
    sources["cell"] = _curve(tmp_path / "cell.csv", [4.0, 5.0], [2.0, 3.0])
    result = IntegratedReTool().run([], _params(include_ir=True), extra=sources)
    assert result.tables["Cell comparison"].empty
    table = plot_tables(result)["Cell comparison"]
    for corrected in (False, True):
        voltage = "Measured (2E) V iR corrected (V)" if corrected else "Measured (2E) V (V)"
        _matches_line(table, voltage, "Measured (2E) |j| (A/cm²)",
                      _line(result, "Cell comparison", "Measured cell", corrected=corrected))
    np.testing.assert_allclose(table["Measured (2E) |j| (A/cm²)"].dropna(), [2.0, 3.0])


def test_optional_cell_does_not_create_empty_measured_columns(tmp_path):
    result = IntegratedReTool().run([], _params(include_ir=True),
                                    extra=_sources(tmp_path, cell=False))
    table = plot_tables(result)["Cell comparison"]
    assert table.columns.tolist() == [
        "Reconstructed V (V)", "Reconstructed |j| (A/cm²)",
        "Reconstructed V iR corrected (V)",
    ]


def test_calibration_exports_alignment_curves_and_feature_markers(tmp_path):
    tool = IntegratedReTool()
    result = tool.run([], _params(correction_mode="3E onset", include_ir=True),
                      extra=_sources(tmp_path, calibration=True))
    tables = plot_tables(result)
    assert list(tables) == ["Cell comparison", "Electrode LSVs", "Calibration alignment"]
    table = tables["Calibration alignment"]
    lines = result.figures["Calibration alignment"].axes[0].lines
    assert len(lines) == 3
    _matches_line(table, "Integrated E iR corrected (V vs RHE)", "Integrated j (A/cm²)", lines[0])
    _matches_line(table, "Integrated E iR + RE corrected (V vs RHE)", "Integrated j (A/cm²)", lines[1])
    _matches_line(table, "3E E iR corrected (V vs RHE)", "3E j (A/cm²)", lines[2])

    workbook = next(artifact for artifact in tool.exportable_artifacts(result)
                    if artifact.kind == "workbook").payload
    sheet = workbook["Calibration alignment"]
    markers = np.asarray(result.figures["Calibration alignment"].axes[0].collections[0].get_offsets())
    numeric_rows = [[cell.value for cell in row if isinstance(cell.value, (int, float))]
                    for row in sheet.iter_rows(min_row=1, max_row=9)]
    assert any(np.allclose(row, markers[0]) for row in numeric_rows if len(row) == 2)
    assert any(np.allclose(row, markers[1]) for row in numeric_rows if len(row) == 2)


def test_workbook_is_readable_and_preserves_numerical_precision(tmp_path):
    tool = IntegratedReTool()
    result = tool.run([], _params(correction_mode="3E onset", include_ir=True,
                                 electrode_axis="Overpotential"),
                      extra=_sources(tmp_path, calibration=True, gaps=True))
    tables = plot_tables(result)
    artifacts = tool.exportable_artifacts(result)
    workbook_artifact = next(artifact for artifact in artifacts if artifact.kind == "workbook")
    assert workbook_artifact.name == "Plot data (Excel)"
    assert workbook_artifact.default_filename == "integrated_re_cell_results.xlsx"
    assert {artifact.name for artifact in artifacts if artifact.kind == "csv"} == {
        "Cell comparison (CSV)", "Electrode LSVs (CSV)",
        "Calibration alignment (CSV)", "Calibration matches (CSV)",
    }
    assert len([artifact for artifact in artifacts if artifact.kind == "png"]) == 3
    assert not any(artifact.kind == "xlsx" for artifact in artifacts)
    target = tmp_path / workbook_artifact.default_filename
    ExportButton._write(workbook_artifact, str(target))
    workbook = load_workbook(target, data_only=False)
    assert workbook.sheetnames == list(tables)
    for name, expected in tables.items():
        sheet = workbook[name]
        header = 10 if name == "Calibration alignment" else 4
        assert sheet.freeze_panes is not None
        assert int("".join(filter(str.isdigit, sheet.freeze_panes))) == header + 1
        assert sheet.sheet_view.showGridLines is False
        assert [cell.value for cell in sheet[header]] == expected.columns.tolist()
        assert all(cell.font.bold for cell in sheet[header])
        assert all(cell.alignment.wrap_text for cell in sheet[header])
        assert all(dimension.width >= 15 for dimension in sheet.column_dimensions.values())
        written = pd.read_excel(target, sheet_name=name, header=header - 1)
        pd.testing.assert_frame_equal(written, expected, check_dtype=False, atol=1e-14, rtol=1e-14)
        for row in sheet.iter_rows(min_row=header + 1):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    assert cell.number_format != "General"
                assert cell.data_type != "f"
        assert sheet.max_row <= header + len(expected)
    assert all("Source" not in column and "HFR" not in column
               and "window" not in column and "drop" not in column
               for table in tables.values() for column in table)


def test_clean_workbook_survives_project_save_and_restore(tmp_path):
    tool = IntegratedReTool()
    result = tool.run([], _params(include_ir=True), extra=_sources(tmp_path))
    artifacts = tool.exportable_artifacts(result)
    snapshot = snapshot_result(result, artifacts)
    frozen = next(item["bytes"] for item in snapshot["artifacts"] if item["kind"] == "workbook")
    save_project(tmp_path / "results.ecwb", {"result": snapshot})
    with load_project(tmp_path / "results.ecwb") as project:
        restored, restored_artifacts = restore_result(project.state["result"])
        workbook = next(item for item in restored_artifacts if item.kind == "workbook")
        target = tmp_path / "restored.xlsx"
        ExportButton._write(workbook, str(target))
        assert target.read_bytes() == frozen
        assert artifact_bytes(workbook) == frozen
        assert load_workbook(BytesIO(frozen)).sheetnames == ["Cell comparison", "Electrode LSVs"]
        pd.testing.assert_frame_equal(restored.tables["EIS spectra"], result.tables["EIS spectra"])


def test_run_warnings_remain_available_separately_from_plot_data(tmp_path):
    tool = IntegratedReTool()
    result = tool.run([], _params(), extra=_sources(tmp_path))
    result.warnings = ["One invalid measurement row is retained as a gap."]
    artifacts = tool.exportable_artifacts(result)
    warning_export = next(item for item in artifacts if item.kind == "txt")
    target = tmp_path / warning_export.default_filename
    ExportButton._write(warning_export, str(target))
    assert target.read_text(encoding="utf-8-sig").strip() == result.warnings[0]
    workbook = next(item.payload for item in artifacts if item.kind == "workbook")
    assert workbook.sheetnames == ["Cell comparison", "Electrode LSVs"]
