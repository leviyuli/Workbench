"""Portable workspace archives, safe failures, and lossless result snapshots."""
from io import BytesIO
import json
import os
from pathlib import Path
import zipfile

import numpy as np
import pandas as pd
from pandas.testing import assert_frame_equal
import pytest
from matplotlib.figure import Figure
from openpyxl import Workbook, load_workbook

from workbench.core import project
from workbench.core.project import (
    ProjectError, artifact_bytes, decode_value, encode_value, load_project,
    restore_result, save_project, snapshot_result,
)
from workbench.core.tool_base import Artifact, ToolResult


def _rewrite_archive(path, change):
    with zipfile.ZipFile(path) as archive:
        data = {name: archive.read(name) for name in archive.namelist()}
    manifest = json.loads(data["manifest.json"])
    change(manifest, data)
    data["manifest.json"] = json.dumps(manifest).encode()
    with zipfile.ZipFile(path, "w") as archive:
        for name, value in data.items():
            archive.writestr(name, value)


def test_archive_embeds_same_named_unicode_inputs_in_order_and_preserves_mtime(tmp_path):
    first = tmp_path / "first" / "LSV 阴极.txt"
    second = tmp_path / "second" / first.name
    first.parent.mkdir()
    second.parent.mkdir()
    first.write_text("first source", encoding="utf-8")
    second.write_text("second source", encoding="utf-8")
    first_time, second_time = 1_700_000_000_100_000_000, 1_700_000_001_200_000_000
    os.utime(first, ns=(first_time, first_time))
    os.utime(second, ns=(second_time, second_time))
    target = tmp_path / "portable.ecwb"
    state = {"pages": {"lsv": {"files": [second, first],
                                      "params": {"selected": first}}},
             "page": "lsv", "annotations": {"shift": "−0.2 V"}}
    save_project(target, state)
    first.unlink()
    second.unlink()
    moved = tmp_path / "moved.ecwb"
    target.rename(moved)

    with load_project(moved) as loaded:
        restored = loaded.state["pages"]["lsv"]
        paths = restored["files"]
        assert [p.read_text(encoding="utf-8") for p in paths] == ["second source", "first source"]
        assert paths[0].name == paths[1].name == first.name
        assert paths[0].parent != paths[1].parent
        assert [p.stat().st_mtime_ns for p in paths] == [second_time, first_time]
        assert restored["params"]["selected"] == paths[1]
        assert loaded.files[str(first)] == paths[1]
        extracted = paths[0]
    assert not extracted.exists()


def test_resave_keeps_original_source_metadata_not_extraction_paths(tmp_path):
    source = tmp_path / "source.txt"
    source.write_text("frequency real imaginary")
    target = tmp_path / "project.ecwb"
    save_project(target, {"input": source})
    with load_project(target) as loaded:
        source.unlink()
        save_project(tmp_path / "resaved.ecwb", loaded.state,
                     source_metadata=loaded.source_metadata)
    with load_project(tmp_path / "resaved.ecwb") as loaded:
        assert set(loaded.files) == {str(source)}
        assert loaded.state["input"].name == "source.txt"


def test_readded_original_remains_distinct_from_its_embedded_copy(tmp_path):
    source = tmp_path / "sample.txt"
    source.write_text("original content")
    target = tmp_path / "project.ecwb"
    save_project(target, {"input": source})
    with load_project(target) as first:
        source.write_text("later edited original")
        save_project(tmp_path / "both.ecwb", {"files": [first.state["input"], source]},
                     source_metadata=first.source_metadata)
    with load_project(tmp_path / "both.ecwb") as second:
        assert [path.read_text() for path in second.state["files"]] == [
            "original content", "later edited original"]


def test_dataframe_structured_json_preserves_mixed_values_precision_and_metadata():
    columns = pd.MultiIndex.from_tuples([("data", "float"), ("data", "integer"),
                                         ("source", "mixed")], names=["group", "quantity"])
    frame = pd.DataFrame({
        "a": [0.12345678901234568, np.inf, -np.inf, np.nan],
        "b": pd.array([2**60 + 1, None, -5, 0], dtype="Int64"),
        "c": pd.Series([None, "", pd.NA, "=literal"], dtype="object"),
    })
    frame.columns = columns
    frame.index = pd.Index(["a", "a", "b", "c"], name="spectrum")
    frame.attrs = {"units": "Ω·cm²", "tuple": (1, 2)}
    encoded = encode_value(frame)
    raw = json.dumps(encoded, allow_nan=False)
    assert "pickle" not in raw
    restored = decode_value(json.loads(raw))
    assert_frame_equal(restored, frame)
    assert restored.attrs == frame.attrs
    assert restored.iloc[0, 2] is None
    assert restored.iloc[2, 2] is pd.NA


@pytest.mark.parametrize("frame", [
    pd.DataFrame(),
    pd.DataFrame(index=pd.RangeIndex(4, 10, 2, name="run")),
    pd.DataFrame([[1, "x"]], columns=["same", "same"]),
    pd.DataFrame({"value": pd.Categorical(["a", None], categories=["a", "b"], ordered=True)}),
    pd.DataFrame({"date": pd.date_range("2026-01-01", periods=2, tz="Asia/Shanghai")},
                 index=pd.date_range("2025-01-01", periods=2, freq="h", name="time")),
    pd.DataFrame({"value": pd.array([True, None], dtype="boolean")}),
    pd.DataFrame({"value": pd.array(["", None], dtype="string")}),
])
def test_dataframe_empty_duplicate_categorical_datetime_and_extension_dtypes(frame):
    assert_frame_equal(decode_value(encode_value(frame)), frame)


def test_numpy_arrays_and_nonfinite_values_are_explicit_json():
    state = {"array": np.array([[1, np.nan], [np.inf, -np.inf]]),
             "scalar": np.int64(2**60), "when": np.datetime64("2026-01-01", "ns"),
             "tuple": (None, -0.0, complex(1, 2)), "bytes": b"\x00\xff"}
    restored = decode_value(json.loads(json.dumps(encode_value(state), allow_nan=False)))
    np.testing.assert_equal(restored["array"], state["array"])
    assert restored["when"] == state["when"]
    assert restored["when"].dtype == state["when"].dtype
    assert restored["scalar"] == state["scalar"]
    assert restored["tuple"] == state["tuple"]
    assert restored["bytes"] == state["bytes"]


def test_result_roundtrip_restores_static_plot_all_fields_and_exact_exports(tmp_path):
    figure = Figure(figsize=(2, 2))
    figure.subplots().plot([0, 1], [0, 2])
    frame = pd.DataFrame({"j": [0.0, 0.1], "V": [1.23, 1.5]})
    book = Workbook()
    book.active["A1"] = "=1+1"
    book.active["B1"] = "00123"
    book.active["C1"] = "=literal"
    book.active["C1"].data_type = "s"
    book.active["B1"].number_format = "@"
    result = ToolResult(figures={"Cell": figure}, tables={"Data": frame},
                        scalars={"HFR": 0.03}, warnings=["source warning"],
                        text_blocks={"Explanation": "RHE"}, param_updates={"shift": -0.2},
                        param_annotations={"shift": "3E onset"},
                        payload={"not serializable": object()},
                        display_tables=False, display_summary=False)
    artifacts = [Artifact("CSV", "csv", "cell.csv", frame),
                 Artifact("XLSX", "xlsx", "tables.xlsx", {"Data": frame}),
                 Artifact("Workbook", "workbook", "formatted.xlsx", book),
                 Artifact("Text", "txt", "parameters.txt", "HFR: 0.03 Ω·cm²"),
                 Artifact("Plot", "png", "plot.png", figure)]
    state = {"result": snapshot_result(result, artifacts), "inputs": {"shift": 0.7}}
    target = tmp_path / "results.ecwb"
    save_project(target, state)
    with load_project(target) as loaded:
        restored, exports = restore_result(loaded.state["result"])
        assert not restored.figures
        assert list(restored.figure_images) == ["Cell"]
        assert restored.figure_images["Cell"].startswith(b"\x89PNG\r\n\x1a\n")
        assert_frame_equal(restored.tables["Data"], frame)
        assert restored.scalars == result.scalars
        assert restored.warnings == result.warnings
        assert restored.param_updates == {"shift": -0.2}
        assert loaded.state["inputs"]["shift"] == 0.7
        assert restored.param_annotations == result.param_annotations
        assert restored.payload == {}
        assert not restored.display_tables and not restored.display_summary
        assert [(a.name, a.kind, a.default_filename) for a in exports] == [
            (a.name, a.kind, a.default_filename) for a in artifacts]
        assert exports[0].serialized_payload == artifact_bytes(artifacts[0])
        assert exports[3].serialized_payload.decode("utf-8-sig") == artifacts[3].payload
        workbook = load_workbook(BytesIO(exports[2].serialized_payload), data_only=False)
        assert workbook.active["A1"].data_type == "f"
        assert workbook.active["B1"].value == "00123"
        assert workbook.active["B1"].number_format == "@"
        assert workbook.active["C1"].data_type == "s"
        workbook.close()
        # Restored results can themselves be saved repeatedly without a rerun.
        second = snapshot_result(restored, exports)
        assert second["figures"] == state["result"]["figures"]
        assert [a["bytes"] for a in second["artifacts"]] == [
            a["bytes"] for a in state["result"]["artifacts"]]


@pytest.mark.parametrize("mode", ["missing", "version", "checksum", "unsafe", "unknown_type"])
def test_invalid_archive_fails_and_cleans_extraction_directory(tmp_path, monkeypatch, mode):
    source = tmp_path / "source.txt"
    source.write_text("sample")
    target = tmp_path / "invalid.ecwb"
    save_project(target, {"input": source})
    if mode == "missing":
        change = lambda manifest, data: data.pop(manifest["files"][0]["member"])
    elif mode == "version":
        change = lambda manifest, data: manifest.update(version=999)
    elif mode == "checksum":
        change = lambda manifest, data: data.update({manifest["files"][0]["member"]: b"wrong!"})
    elif mode == "unsafe":
        change = lambda manifest, data: data.update({"../outside.txt": b"malicious"})
    else:
        change = lambda manifest, data: manifest.update(state={"$ecwb": "python_pickle"})
    _rewrite_archive(target, change)
    created = []
    original_temporary = project.tempfile.TemporaryDirectory

    def create(*args, **kwargs):
        directory = original_temporary(*args, **kwargs)
        created.append(Path(directory.name))
        return directory

    monkeypatch.setattr(project.tempfile, "TemporaryDirectory", create)
    with pytest.raises(ProjectError):
        load_project(target)
    assert created and all(not folder.exists() for folder in created)
    assert not (tmp_path / "outside.txt").exists()


def test_failed_save_preserves_existing_target_and_removes_temporary_archive(tmp_path, monkeypatch):
    target = tmp_path / "important.ecwb"
    save_project(target, {"value": 1})
    original = target.read_bytes()
    with pytest.raises(ProjectError, match="unavailable"):
        save_project(target, {"input": tmp_path / "missing.txt"})
    assert target.read_bytes() == original

    def fail_replace(*_):
        raise OSError("simulated full or locked target")

    monkeypatch.setattr(project.os, "replace", fail_replace)
    with pytest.raises(ProjectError, match="simulated"):
        save_project(target, {"value": 2})
    assert target.read_bytes() == original
    assert not list(tmp_path.glob(".ecwb-*.tmp"))


def test_temporary_archive_creation_failure_is_a_project_error(tmp_path, monkeypatch):
    target = tmp_path / "important.ecwb"
    save_project(target, {"value": 1})
    original = target.read_bytes()

    def fail_create(*args, **kwargs):
        raise OSError("Disk full")

    monkeypatch.setattr(project.tempfile, "mkstemp", fail_create)
    with pytest.raises(ProjectError, match="Could not save project: Disk full"):
        save_project(target, {"value": 2})
    assert target.read_bytes() == original
    assert not list(tmp_path.glob(".ecwb-*.tmp"))


def test_snapshot_rejects_invalid_png_and_unknown_export_kind():
    result = ToolResult(figure_images={"Broken": b"not a png"})
    with pytest.raises(ProjectError, match="Invalid saved result"):
        restore_result(snapshot_result(result, []))
    with pytest.raises(ProjectError, match="Unknown export kind"):
        artifact_bytes(Artifact("Unsafe", "pickle", "data.pkl", object()))


def test_blob_integrity_and_duplicate_member_validation(tmp_path):
    target = tmp_path / "bytes.ecwb"
    save_project(target, {"image": b"binary bytes"})
    with zipfile.ZipFile(target) as archive:
        blob = next(name for name in archive.namelist() if name.startswith("objects/"))
    _rewrite_archive(target, lambda manifest, data: data.update({blob: b"wrong bytes!"}))
    with pytest.raises(ProjectError, match="checksum"):
        load_project(target)
    save_project(target, {})
    with pytest.warns(UserWarning, match="Duplicate"):
        with zipfile.ZipFile(target, "a") as archive:
            archive.writestr("manifest.json", "{}")
    with pytest.raises(ProjectError, match="duplicate"):
        load_project(target)
