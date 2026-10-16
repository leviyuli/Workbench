"""Workbook formatting must leave scientific values and source records intact."""

from datetime import datetime

import pandas as pd
import pytest
from openpyxl import load_workbook

from workbench.core.io import write_xlsx
from workbench.tools.txt_xlsx.logic import convert_to_xlsx


def test_analysis_workbook_retains_precision_gaps_and_cell_types(tmp_path):
    path = tmp_path / "results.xlsx"
    frame = pd.DataFrame({
        "Potential (V vs RHE)": [0.123456789012345, -0.987654321, None],
        "Current density (A/cm²)": [2.34e-9, 0.0, 1.23e8],
        "Selected": [True, False, True],
        "Measured at": [datetime(2026, 10, 16, 12), None, None],
    })
    write_xlsx({"Curves": frame}, path)

    with pd.ExcelFile(path) as book:
        restored = pd.read_excel(book, sheet_name="Curves")
    assert restored["Potential (V vs RHE)"].tolist()[:2] == pytest.approx(
        frame["Potential (V vs RHE)"].tolist()[:2], abs=1e-15
    )
    assert pd.isna(restored.iloc[2, 0])
    workbook = load_workbook(path)
    sheet = workbook["Curves"]
    assert sheet["B2"].value == pytest.approx(2.34e-9)
    assert sheet["B3"].value == 0
    assert "E+00" in sheet["B2"].number_format
    assert "E+00" in sheet["B4"].number_format
    assert sheet["C2"].data_type == "b"
    assert sheet["D2"].value == datetime(2026, 10, 16, 12)
    assert sheet.freeze_panes == "A2"
    assert sheet.sheet_view.showGridLines is True
    assert sheet["A1"].font.bold and sheet["A1"].alignment.wrap_text
    for row in sheet:
        for cell in row:
            if cell.value is not None:
                assert cell.font.name == "Calibri"
                assert cell.font.sz == 11
                assert cell.font.color.rgb[-6:] == "000000"
                assert cell.fill.patternType is None
                assert all(getattr(getattr(cell.border, side), "style", None) is None
                           for side in ("left", "right", "top", "bottom"))
    workbook.close()


def test_long_headers_and_multiline_results_wrap_without_expanding_curve_columns(tmp_path):
    path = tmp_path / "labels.xlsx"
    label = "Reconstructed cell voltage after correction (V)"
    write_xlsx({"Curve": pd.DataFrame({label: [1.5], "Condition": ["Hg/HgO\n1 M KOH"]})}, path)
    workbook = load_workbook(path)
    sheet = workbook.active
    assert sheet.column_dimensions["A"].width <= 34
    assert sheet.row_dimensions[1].height == 30
    assert sheet["B2"].alignment.wrap_text
    assert sheet.row_dimensions[2].height == 30
    workbook.close()


def test_explicit_analysis_number_format_is_honored(tmp_path):
    path = tmp_path / "format.xlsx"
    write_xlsx({"Values": pd.DataFrame({"x": [1, 1.23456, 1e-9]})}, path, number_format="0.0000")
    workbook = load_workbook(path)
    assert [workbook.active.cell(row, 1).number_format for row in (2, 3, 4)] == ["0.0000"] * 3
    workbook.close()


def test_converter_formats_headers_without_changing_data_or_requested_format(tmp_path):
    path = tmp_path / "source.csv"
    path.write_text("Identifier,Potential (V),Note\n00123,0.123456789,first\n00042,-2,second\n", encoding="utf-8")
    result = convert_to_xlsx([path], number_format="0.00000000")
    output = tmp_path / "converted.xlsx"
    result["workbook"].save(output)
    workbook = load_workbook(output)
    sheet = workbook.active
    assert sheet.max_row == 3
    assert sheet.freeze_panes == "A2"
    assert sheet["A1"].font.bold
    assert sheet["A2"].value == "00123" and sheet["A2"].data_type == "s"
    assert sheet["B2"].value == 0.123456789
    assert sheet["B2"].number_format == "0.00000000"
    workbook.close()


def test_converter_numeric_first_row_remains_data(tmp_path):
    path = tmp_path / "numeric.dat"
    path.write_text("0\t1.23456\n1\t2.34567\n", encoding="utf-8")
    result = convert_to_xlsx([path])
    sheet = result["workbook"].active
    assert list(sheet.values) == [(0, 1.23456), (1, 2.34567)]
    assert sheet.freeze_panes is None
    assert not sheet["A1"].font.bold
    assert sheet["B1"].number_format == "General"


def test_converter_default_does_not_display_small_scientific_values_as_zero(tmp_path):
    from workbench.tools.txt_xlsx.tool import TxtXlsxTool

    path = tmp_path / "small.csv"
    path.write_text("Current (A)\n0.000000123\n", encoding="utf-8")
    tool = TxtXlsxTool()
    params = {spec.key: spec.default for spec in tool.parameter_schema()}
    result = tool.run([path], params)
    workbook = tool.exportable_artifacts(result)[0].payload
    assert workbook.active["A2"].value == pytest.approx(1.23e-7)
    assert workbook.active["A2"].number_format == "General"
