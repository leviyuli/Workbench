"""Regression tests for shared I/O and TXT/XLSX conversion hardening."""

from __future__ import annotations

import csv

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from workbench.core import io as core_io
from workbench.tools.txt_xlsx.logic import convert_to_xlsx


def test_read_text_grid_uses_real_csv_parser(tmp_path):
    source = tmp_path / "quoted.csv"
    source.write_text(
        'name,id,note\n"Doe, Jane","00123","first line\nsecond line"\n',
        encoding="utf-8",
    )

    grid = core_io.read_text_grid(source)

    assert grid == [
        ["name", "id", "note"],
        ["Doe, Jane", "00123", "first line\nsecond line"],
    ]


def test_read_text_grid_collapses_repeated_whitespace(tmp_path):
    source = tmp_path / "spaced.txt"
    source.write_text("time    current\n0       1.5\n", encoding="utf-8")

    assert core_io.read_text_grid(source) == [
        ["time", "current"],
        ["0", "1.5"],
    ]


def test_read_table_collapses_repeated_whitespace(tmp_path):
    source = tmp_path / "spaced.txt"
    source.write_text("time    current\n0       1.5\n", encoding="utf-8")

    frame = core_io.read_table(source)

    assert list(frame.columns) == ["time", "current"]
    assert frame.to_dict("records") == [{"time": 0, "current": 1.5}]


def test_dataframe_exports_neutralize_formula_text(tmp_path):
    frame = pd.DataFrame({
        "=header": ["=1+1", "+SUM(A1:A2)", "-cmd|' /C calc'!A0", "@SUM(1,1)", -2],
    })

    csv_path = tmp_path / "safe.csv"
    core_io.write_csv(frame, csv_path)
    with csv_path.open(newline="", encoding="utf-8-sig") as handle:
        rows = list(csv.reader(handle))
    assert rows[0] == ["'=header"]
    assert [row[0] for row in rows[1:5]] == [
        "'=1+1",
        "'+SUM(A1:A2)",
        "'-cmd|' /C calc'!A0",
        "'@SUM(1,1)",
    ]
    assert rows[5] == ["-2"]  # a genuine numeric value remains numeric

    xlsx_path = tmp_path / "safe.xlsx"
    core_io.write_xlsx({"Results": frame}, xlsx_path)
    workbook = load_workbook(xlsx_path, data_only=False)
    sheet = workbook["Results"]
    assert sheet["A1"].value == "'=header"
    assert sheet["A2"].value == "'=1+1"
    assert sheet["A2"].data_type == "s"
    assert sheet["A6"].value == -2
    workbook.close()


def test_text_exports_mark_utf8_and_preserve_scientific_units(tmp_path):
    frame = pd.DataFrame({
        "Quantity": ["Series HFR", "Terminal current magnitude"],
        "Value": [0.02, 10],
        "Unit": ["Ω·cm²", "A/cm²"],
    })

    csv_path = tmp_path / "units.csv"
    core_io.write_csv(frame, csv_path)
    assert csv_path.read_bytes().startswith(b"\xef\xbb\xbf")
    assert pd.read_csv(csv_path, encoding="utf-8-sig").equals(frame)

    txt_path = tmp_path / "units.txt"
    text = "Series HFR\t0.02\tΩ·cm²\nTerminal current magnitude\t10\tA/cm²\n"
    core_io.write_txt(text, txt_path)
    assert txt_path.read_bytes().startswith(b"\xef\xbb\xbf")
    assert txt_path.read_text(encoding="utf-8-sig") == text


def test_write_xlsx_sanitizes_titles_and_casefold_collisions(tmp_path):
    path = tmp_path / "titles.xlsx"
    core_io.write_xlsx(
        {
            "Bad/Name": pd.DataFrame({"x": [1]}),
            "bad:name": pd.DataFrame({"x": [2]}),
            "'": pd.DataFrame({"x": [3]}),
            "x" * 40: pd.DataFrame({"x": [4]}),
        },
        path,
    )

    workbook = load_workbook(path, read_only=True)
    names = workbook.sheetnames
    assert names[0] == "Bad_Name"
    assert names[1] == "bad_name_1"
    assert names[2] == "Sheet"
    assert len(names[3]) == 31
    assert len({name.casefold() for name in names}) == len(names)
    workbook.close()


def test_converter_preserves_identifiers_and_converts_safe_numbers(tmp_path):
    source = tmp_path / "identifiers.csv"
    source.write_text(
        "kind,value\n"
        "leading,00123\n"
        "leading_decimal,001.25\n"
        "long,1234567890123456\n"
        "integer,42\n"
        "decimal,3.5\n"
        "negative,-3\n"
        "payload,=2+2\n",
        encoding="utf-8",
    )

    result = convert_to_xlsx([source])
    sheet = result["workbook"]["identifiers"]

    assert sheet["B2"].value == "00123"
    assert sheet["B2"].data_type == "s"
    assert sheet["B3"].value == "001.25"
    assert sheet["B3"].data_type == "s"
    assert sheet["B4"].value == "1234567890123456"
    assert sheet["B4"].data_type == "s"
    assert sheet["B5"].value == 42
    assert sheet["B6"].value == 3.5
    assert sheet["B7"].value == -3
    assert sheet["B8"].value == "'=2+2"
    assert sheet["B8"].data_type == "s"


def test_converter_retains_xlsx_formulas_but_not_literal_formula_text(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "=1+1"
    sheet["A2"] = "ordinary"
    sheet["A3"] = "=literal text"
    sheet["A3"].data_type = "s"
    workbook.save(source)
    workbook.close()

    result = convert_to_xlsx([source])
    copied = result["workbook"]["Data"]

    assert copied["A1"].value == "=1+1"
    assert copied["A1"].data_type == "f"
    assert copied["A2"].value == "ordinary"
    assert copied["A2"].data_type == "s"
    assert copied["A3"].value == "=literal text"
    assert copied["A3"].data_type == "s"

    output = tmp_path / "copied.xlsx"
    result["workbook"].save(output)
    reopened = load_workbook(output, data_only=False)
    assert reopened["Data"]["A1"].data_type == "f"
    assert reopened["Data"]["A1"].value == "=1+1"
    assert reopened["Data"]["A3"].data_type == "s"
    assert reopened["Data"]["A3"].value == "=literal text"
    reopened.close()


def test_converter_deduplicates_sheet_names_case_insensitively(tmp_path):
    paths = []
    for filename, title in (("a.xlsx", "Data"), ("b.xlsx", "data")):
        path = tmp_path / filename
        workbook = Workbook()
        workbook.active.title = title
        workbook.active["A1"] = 1
        workbook.save(path)
        workbook.close()
        paths.append(path)

    result = convert_to_xlsx(paths)

    assert result["sheets"] == ["Data", "data_1"]


def test_text_size_and_cell_limits_fail_cleanly(tmp_path, monkeypatch):
    source = tmp_path / "large.csv"
    source.write_text("a,b\n1,2\n", encoding="utf-8")

    monkeypatch.setattr(core_io, "MAX_TEXT_BYTES", 4)
    with pytest.raises(ValueError, match="too large"):
        core_io.read_text_grid(source)

    monkeypatch.setattr(core_io, "MAX_TEXT_BYTES", 1024)
    monkeypatch.setattr(core_io, "MAX_TABLE_CELLS", 3)
    with pytest.raises(ValueError, match="too many cells"):
        core_io.read_text_grid(source)


def test_xlsx_decompression_limit_and_invalid_archive_fail_cleanly(tmp_path, monkeypatch):
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active["A1"] = "value"
    workbook.save(source)
    workbook.close()

    monkeypatch.setattr(core_io, "MAX_XLSX_UNCOMPRESSED_BYTES", 1)
    with pytest.raises(ValueError, match="expands to"):
        core_io.read_xlsx_grids(source)

    invalid = tmp_path / "invalid.xlsx"
    invalid.write_text("not a zip archive", encoding="utf-8")
    monkeypatch.setattr(core_io, "MAX_XLSX_UNCOMPRESSED_BYTES", 1024 * 1024)
    with pytest.raises(ValueError, match="not a valid .xlsx"):
        core_io.read_xlsx_grids(invalid)


def test_xls_missing_dependency_error_is_actionable(tmp_path, monkeypatch):
    source = tmp_path / "legacy.xls"
    source.write_bytes(b"placeholder")
    real_find_spec = core_io.importlib.util.find_spec
    monkeypatch.setattr(
        core_io.importlib.util,
        "find_spec",
        lambda name: None if name == "xlrd" else real_find_spec(name),
    )

    with pytest.raises(ValueError, match=r"xlrd>=2\.0\.1"):
        core_io.read_excel(source)
