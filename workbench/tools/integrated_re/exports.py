"""Final plotted series and their Excel export."""
from __future__ import annotations

import numpy as np
import pandas as pd
from openpyxl import Workbook

from workbench.core.excel import format_sheet
from workbench.core.tool_base import ToolResult


CURRENT = "Current density (A/cm²)"


def _column(table: pd.DataFrame, key: str, mask=None) -> pd.Series:
    values = table[key]
    return (values if mask is None else values.where(mask)).reset_index(drop=True)


def plot_tables(result: ToolResult) -> dict[str, pd.DataFrame]:
    """Preserve each plotted trace's sampling order, current grid and gaps."""
    tables = result.tables
    reconstructed = tables["Reconstructed cell"]
    cell = {
        "Reconstructed V (V)": _column(reconstructed, "Reconstructed cell voltage (V)"),
        "Reconstructed |j| (A/cm²)": _column(reconstructed, CURRENT),
    }
    corrected = tables.get("Reconstructed cell (iR corrected)")
    if corrected is not None:
        cell["Reconstructed V iR corrected (V)"] = _column(corrected, "Reconstructed cell voltage (V)")
    measured = tables.get("Cell LSV")
    if measured is not None:
        branch = measured[CURRENT] >= 0
        cell["Measured (2E) V (V)"] = _column(measured, "Cell voltage (V)", branch)
        cell["Measured (2E) |j| (A/cm²)"] = _column(measured, CURRENT, branch)
        if corrected is not None:
            cell["Measured (2E) V iR corrected (V)"] = _column(measured, "iR-corrected cell voltage (V)", branch)

    settings = dict(tables["Potential settings"].itertuples(index=False, name=None))
    overpotential = settings["Electrode plot scale"] == "Overpotential"
    value_key = "Overpotential (V)" if overpotential else "E (V vs RHE)"
    corrected_key = "iR-corrected overpotential (V)" if overpotential else "iR-corrected E (V vs RHE)"
    label = "η (V)" if overpotential else "E (V vs RHE)"
    corrected_label = "η iR corrected (V)" if overpotential else "E iR corrected (V vs RHE)"
    electrodes = {}
    for role in ("Anode", "Cathode"):
        curve = tables[f"{role} LSV"]
        usable = np.isfinite(curve[value_key]) & np.isfinite(curve[CURRENT])
        electrodes[f"{role} {label}"] = _column(curve, value_key, usable)
        electrodes[f"{role} j (A/cm²)"] = _column(curve, CURRENT, usable)
        if corrected is not None:
            electrodes[f"{role} {corrected_label}"] = _column(curve, corrected_key, usable)

    output = {"Cell comparison": pd.DataFrame(cell), "Electrode LSVs": pd.DataFrame(electrodes)}
    if "Calibration LSV" in tables:
        integrated = tables["Calibration integrated LSV"]
        calibration = tables["Calibration LSV"]
        output["Calibration alignment"] = pd.DataFrame({
            "Integrated E iR corrected (V vs RHE)": _column(integrated, "E before correction (V vs RHE)"),
            "Integrated E iR + RE corrected (V vs RHE)": _column(integrated, "E (V vs RHE)"),
            "Integrated j (A/cm²)": _column(integrated, CURRENT),
            "3E E iR corrected (V vs RHE)": _column(calibration, "E (V vs RHE)"),
            "3E j (A/cm²)": _column(calibration, CURRENT),
        })
    return output


def _write_table(sheet, frame: pd.DataFrame, header_row: int) -> None:
    for column, heading in enumerate(frame.columns, start=1):
        sheet.cell(header_row, column, str(heading))
    for row_number, values in enumerate(frame.itertuples(index=False, name=None), start=header_row + 1):
        for column, value in enumerate(values, start=1):
            if pd.isna(value) or isinstance(value, (float, np.floating)) and not np.isfinite(value):
                continue
            if isinstance(value, np.generic):
                value = value.item()
            sheet.cell(row_number, column, value)


def results_workbook(result: ToolResult, tables: dict[str, pd.DataFrame]) -> Workbook:
    """One sheet per final plot, with calibration markers above their curves."""
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, frame in tables.items():
        sheet = workbook.create_sheet(name)
        header_row = 1
        if name == "Calibration alignment":
            features = result.tables["Calibration features"][[
                "Source", "E after correction (V vs RHE)", CURRENT,
            ]].rename(columns={"E after correction (V vs RHE)": "E (V vs RHE)", CURRENT: "j (A/cm²)"})
            _write_table(sheet, features, 1)
            header_row = 5
        _write_table(sheet, frame, header_row)
        format_sheet(sheet, header_row=header_row,
                     header_rows=(1, 5) if header_row == 5 else (1,))
    return workbook
